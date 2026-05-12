import os 
import json
from PIL import Image
from PIL.ExifTags import TAGS
from PyPDF2 import PdfReader
from docx import Document 
from openpyxl import load_workbook


def get_image_meta(path):
    meta = {}
    try: 
        with Image.open(path) as imag:
            exif = img._getexif()
            if exif:
                for tag, value in exif.items():
                    decoded = TAGS.get(tag, tag)
                    meta[decoded] = str(value)


    except : pass
    return meta


def get_pdf_meta(path):
    meta = {}
    try: 
        reader = PdfReader(path)
        info = reader.metadata
        for key in info: 
            meta [key.replace('/','')] = str(info[key])

    except: pass
    return meta



def get_office_meta(path):
    meta = {}
    try: 
        if path.endswith('.docx'):
            doc = Document(path)
            prop = doc.core_properties

        elif path.endswith('.xlsx'):
            wb = load_workbook(path)
            prop = wb.properties

        attrs = ['author', 'created', 'last_modified_by','last_printed','title','version']
        for attr in attrs:
            if hasattr(prop, attr):
                meta[attr] = str(getattr(prop, attr))

    except: pass
    return meta


def main():
    print("METATRACTOR UNIVERSAL - OSINT Tool")
    folder = input("Introduce la ruta de la carpeta: ").strip('" ')#limpiar comillas de windows

    if not os.path.exists(folder):
        print("Ruta no valida")
        return
    
    report = {}
    for root, _, files in os.walks(folder):
        for file in files:
            path = os.path.join(root, file)
            ext = file.lower()

            if ext.endswith(('.jpg', '.jpeg','.png')):
                report[file] = get_image_meta(path)
            elif ext.endswith('.pdf'):
                report[file] = get_pdf_meta(path)
            elif ext.endswith(('.docx', '.xlsx')):
                report[file] = get_office_meta(path)

        output = os.path.join(os.path.expanduser("-"), "Desktop", "informe-metadatos.json")
        with open(output, "w", indent=4)

        print("Analisis listo. Informe en: {output}")


if __name__== "__main__":
    main()

